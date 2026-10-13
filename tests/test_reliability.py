from __future__ import annotations

import gzip
import io
import subprocess
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from threading import Event

import pymupdf as fitz
import pytest
from docx import Document
from PIL import Image

import app as service
from converter import ConversionError, ConversionManager, output_filename
from lifecycle import BrowserSessionManager


@pytest.fixture(autouse=True)
def clear_jobs():
    yield
    service.cleanup_all_jobs()


def post_texts(client, *, headers=None, files=None):
    names = files or ["notes.txt"]
    return client.post(
        "/api/convert",
        data={
            "files": [(io.BytesIO(b"hello"), name) for name in names],
            "targets": ["txt"] * len(names),
        },
        headers=headers or {},
    )


def test_progress_and_busy_response_during_conversion(monkeypatch):
    started, finish = Event(), Event()
    original_convert = service.converter.convert
    job_id = uuid.uuid4().hex

    def delayed(*args):
        started.set()
        assert finish.wait(5)
        return original_convert(*args)

    monkeypatch.setattr(service.converter, "convert", delayed)
    with ThreadPoolExecutor(max_workers=1) as executor:
        pending = executor.submit(
            post_texts, service.app.test_client(), headers={"X-Conversion-ID": job_id}
        )
        try:
            assert started.wait(5)
            with service.app.test_client() as client:
                status = client.get(f"/api/jobs/{job_id}").get_json()
                assert status["status"] == "running"
                assert status["current_file"] == "notes.txt"
                assert status["completed"] == 0
                busy = post_texts(client)
                assert busy.status_code == 429
                assert busy.headers["Retry-After"] == "3"
                assert client.delete(f"/api/jobs/{job_id}").status_code == 409
                service.JOBS[job_id]["created"] = 0
                service.cleanup_expired_jobs()
                assert job_id in service.JOBS
        finally:
            finish.set()
        response = pending.result(timeout=5)
    assert response.status_code == 200
    assert response.get_json()["completed"] == 1
    assert not list(service.JOBS[job_id]["directory"].glob("source-*"))


def test_download_survives_clear_and_cleanup():
    with service.app.test_client() as client:
        payload = post_texts(client).get_json()
        job_id = payload["job_id"]
        directory = service.JOBS[job_id]["directory"]
        file_id = payload["results"][0]["file_id"]
        response = client.get(f"/api/jobs/{job_id}/files/{file_id}", buffered=False)
        try:
            assert response.status_code == 200
            assert client.delete(f"/api/jobs/{job_id}").status_code == 204
            assert client.get(f"/api/jobs/{job_id}").status_code == 404
            assert directory.exists()
            assert response.data == b"hello"
        finally:
            response.close()
        assert not directory.exists()
        assert job_id not in service.JOBS


def test_concurrent_zip_downloads_are_complete():
    with service.app.test_client() as client:
        payload = post_texts(client, files=["same.txt", "SAME.txt"]).get_json()

    def download():
        with (
            service.app.test_client() as client,
            client.get(payload["download_all_url"]) as response,
        ):
            assert response.status_code == 200
            return response.data

    with ThreadPoolExecutor(max_workers=4) as executor:
        archives = list(executor.map(lambda _: download(), range(4)))
    for data in archives:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert archive.testzip() is None
            assert archive.namelist() == ["same.txt", "SAME (2).txt"]
    assert not list(service.JOBS[payload["job_id"]]["directory"].glob("*.part"))


def test_partial_failure_does_not_reserve_name_or_keep_sources():
    with service.app.test_client() as client:
        response = client.post(
            "/api/convert",
            data={
                "files": [(io.BytesIO(b""), "notes.txt"), (io.BytesIO(b"valid"), "notes.txt")],
                "targets": ["txt", "txt"],
            },
        )
        payload = response.get_json()
        assert payload["success_count"] == 1
        assert payload["results"][0]["status"] == "error"
        assert payload["results"][1]["output_name"] == "notes.txt"
        assert len(list(service.JOBS[payload["job_id"]]["directory"].iterdir())) == 1


def test_failed_batch_and_invalid_requests_are_safe(monkeypatch):
    with service.app.test_client() as client:
        assert client.post("/api/convert").status_code == 400
        assert (
            client.post("/api/convert", data={"files": (io.BytesIO(b"x"), "x.txt")}).status_code
            == 400
        )
        assert post_texts(client, headers={"X-Conversion-ID": "bad"}).status_code == 400
        failed = post_texts(client, files=["unknown.abc"]).get_json()
        assert failed["success_count"] == 0
        assert not service.JOBS[failed["job_id"]]["directory"].exists()
        assert client.get(f"/api/jobs/{failed['job_id']}/download-all").status_code == 404
        assert client.get("/api/jobs/absent").status_code == 404
        good = post_texts(client).get_json()
        assert post_texts(client, headers={"X-Conversion-ID": good["job_id"]}).status_code == 409
        missing = client.get(f"/api/jobs/{good['job_id']}/files/absent")
        assert missing.status_code == 404
        assert service.JOBS[good["job_id"]]["downloads"] == 0
        service.JOBS[good["job_id"]]["created"] = 0
        service.cleanup_expired_jobs()
        assert good["job_id"] not in service.JOBS
        monkeypatch.setitem(service.app.config, "MAX_CONTENT_LENGTH", 1)
        assert post_texts(client).status_code == 413


def test_local_origin_and_host_protection():
    with service.app.test_client() as client:
        assert post_texts(client, headers={"Origin": "https://example.com"}).status_code == 403
        assert (
            client.get("/api/capabilities", headers={"Host": "attacker.example"}).status_code == 400
        )
        response = client.get("/api/capabilities")
        assert response.headers["Cache-Control"] == "no-store"
        assert ".jpeg" in response.get_json()["accept"]


@pytest.mark.parametrize("name", ["CON.txt", "lpt1.txt", "aux.report.txt"])
def test_windows_reserved_names_are_safe(name):
    assert output_filename(name, "pdf", set()).startswith("_")
    assert len(output_filename("a" * 255 + ".txt", "pdf", set())) < 200
    assert len(output_filename("🎉" * 255 + ".txt", "pdf", set()).encode("utf-16-le")) < 400


def test_image_orientation_and_single_frame_export(monkeypatch, tmp_path):
    manager = ConversionManager()
    source, output = tmp_path / "phone.jpg", tmp_path / "phone.png"
    image = Image.new("RGB", (12, 8), "red")
    exif = image.getexif()
    exif[274] = 6
    image.save(source, exif=exif)
    manager.convert(source, "jpg", output, "png")
    with Image.open(output) as converted:
        assert converted.size == (8, 12)
        assert converted.getexif().get(274, 1) == 1

    gif = tmp_path / "animated.gif"
    frames = [Image.new("RGB", (12, 8), color) for color in ("red", "blue")]
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=[100, 200], loop=0)
    monkeypatch.setattr(
        "converter.ImageSequence.Iterator", lambda _: pytest.fail("Decoded unused frames")
    )
    assert "first frame" in manager.convert(gif, "gif", output, "png")


def test_docx_content_order(tmp_path):
    source = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Middle"
    document.add_paragraph("After")
    document.save(source)
    manager = ConversionManager()
    for target in ("txt", "html", "pdf"):
        output = tmp_path / f"ordered.{target}"
        manager.convert(source, "docx", output, target)
        if target == "pdf":
            with fitz.open(output) as pdf:
                text = "".join(page.get_text() for page in pdf)
        else:
            text = output.read_text(encoding="utf-8")
        assert text.index("Before") < text.index("Middle") < text.index("After")


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252"])
def test_text_encodings(encoding, tmp_path):
    source = tmp_path / "text.txt"
    source.write_bytes("Café text".encode(encoding))
    assert ConversionManager._read_text(source) == "Café text"


def test_scanned_pdf_has_actionable_error(tmp_path):
    source = tmp_path / "scanned.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(source)
    with pytest.raises(ConversionError, match="OCR"):
        ConversionManager().convert(source, "pdf", tmp_path / "out.txt", "txt")


def test_lottie_expansion_is_bounded(monkeypatch, tmp_path):
    source = tmp_path / "huge.tgs"
    source.write_bytes(gzip.compress(b" " * 2000))
    monkeypatch.setattr("converter.MAX_LOTTIE_BYTES", 1000)
    with pytest.raises(ConversionError, match="safety limit"):
        ConversionManager().convert(source, "tgs", tmp_path / "out.json", "json")


def test_office_timeout_and_private_profile(monkeypatch, tmp_path):
    manager = ConversionManager()
    manager.libreoffice = "soffice"

    def timeout(command, **_kwargs):
        assert command[1].startswith("-env:UserInstallation=file:")
        raise subprocess.TimeoutExpired(command, 600)

    monkeypatch.setattr("converter.subprocess.run", timeout)
    with pytest.raises(ConversionError, match="ten-minute"):
        manager._convert_with_libreoffice(tmp_path / "x.doc", tmp_path / "x.pdf", "pdf")


def test_browser_close_waits_for_active_work():
    now = [0.0]
    shutdowns = []
    manager = BrowserSessionManager(clock=lambda: now[0], shutdown_grace=2, poll_interval=60)
    manager.start(lambda: shutdowns.append(True))
    try:
        token = manager.register()
        manager.begin_activity()
        manager.close(token)
        now[0] = 10
        assert not manager.poll()
        manager.end_activity()
        assert not manager.poll()
        now[0] = 12
        assert manager.poll()
        assert shutdowns == [True]
    finally:
        manager.stop()
